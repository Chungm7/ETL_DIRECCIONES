"""Script oficial para extraer y sincronizar los catálogos de vías metropolitanas y zonas
directamente desde los archivos Excel oficiales de la Municipalidad Provincial de Chiclayo (MPCH).

Genera los diccionarios canónicos con todas sus variantes de abreviaturas, conjugaciones,
variaciones sin tildes y denominaciones comunes para potenciar el matching determinista y de IA.
"""

import json
import logging
import re
import unicodedata
from pathlib import Path
import openpyxl

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("extract_official_catalogs")

BASE_DIR = Path(__file__).resolve().parent.parent
DOCS_DIR = BASE_DIR / "docs" / "cod_vias_y_habilitaciones_urbanas"
CATALOGS_DIR = BASE_DIR / "src" / "catalogs"

EXCEL_VIAS = DOCS_DIR / "CODIFICADOR DE VIAS - CHICLAYO _ FINAL 26-02-2024 (1).xlsx"
EXCEL_ZONAS = DOCS_DIR / "CODIFICADOR DE HABILITACIONES URBANAS -.xlsx"

OUTPUT_VIAS_JSON = CATALOGS_DIR / "vias_chiclayo.json"
OUTPUT_ZONAS_JSON = CATALOGS_DIR / "zonas_chiclayo.json"


def remove_accents(text: str) -> str:
    """Elimina acentos/diacríticos manteniendo la letra base."""
    nfkd = unicodedata.normalize("NFKD", text)
    return "".join([c for c in nfkd if not unicodedata.combining(c)])


def clean_str(val) -> str:
    """Normaliza texto o valores numéricos leídos de celdas de Excel."""
    if val is None:
        return ""
    if isinstance(val, float):
        return str(int(val))
    return str(val).strip()


number_words_map = {
    "1 DE ": ["1 DE ", "UNO DE ", "PRIMERO DE "],
    "2 DE ": ["2 DE ", "DOS DE "],
    "7 DE ENERO": ["7 DE ENERO", "SIETE DE ENERO"],
    "8 DE OCTUBRE": ["8 DE OCTUBRE", "OCHO DE OCTUBRE"],
    "9 DE OCTUBRE": ["9 DE OCTUBRE", "NUEVE DE OCTUBRE"],
    "12 DE OCTUBRE": ["12 DE OCTUBRE", "DOCE DE OCTUBRE"],
    "16 DE OCTUBRE": ["16 DE OCTUBRE", "DIECISEIS DE OCTUBRE"],
    "28 DE JULIO": ["28 DE JULIO", "VEINTIOCHO DE JULIO"],
    "30 DE AGOSTO": ["30 DE AGOSTO", "TREINTA DE AGOSTO"],
    "27 DE NOVIEMBRE": ["27 DE NOVIEMBRE", "VEINTISIETE DE NOVIEMBRE"],
    "4 DE NOVIEMBRE": ["4 DE NOVIEMBRE", "CUATRO DE NOVIEMBRE"],
}


def extract_vias_oficiales():
    """Extrae las 2,935 vías oficiales completas de la hoja 1.FORMATO CODIFICADOR DE VIAS."""
    tipo_via_map = {
        "AV.": 1,
        "CA.": 2,
        "JR.": 3,
        "PJE.": 4,
        "CTRA.": 6,
        "PRLG.": 7,
        "PRLG.AV": 7,
        "PRLG. AV.": 7,
        "PRLG. CA.": 7,
    }

    via_prefixes = {
        "AV.": ["AV.", "AV", "AVENIDA"],
        "CA.": ["CA.", "CA", "CALLE", "CL.", "CL"],
        "JR.": ["JR.", "JR", "JIRON", "JIRÓN"],
        "PJE.": ["PJE.", "PJE", "PASAJE", "PSJ.", "PSJ"],
        "CTRA.": ["CTRA.", "CTRA", "CARRETERA"],
        "PRLG.": ["PRLG.", "PRLG", "PROL.", "PROLONGACION", "PROLONGACIÓN"],
        "PRLG.AV": ["PRLG. AV.", "PRLG AV", "PROL. AV.", "PROLONGACION AVENIDA", "AV.", "AV"],
        "PRLG. AV.": ["PRLG. AV.", "PRLG AV", "PROL. AV.", "PROLONGACION AVENIDA", "AV.", "AV"],
        "PRLG. CA.": ["PRLG. CA.", "PRLG CA", "PROL. CA.", "PROLONGACION CALLE", "CA.", "CA"],
    }

    logger.info("Cargando libro de vías: %s", EXCEL_VIAS.name)
    wb = openpyxl.load_workbook(str(EXCEL_VIAS), data_only=True)
    ws = wb["1.FORMATO CODIFICADOR DE VIAS"]

    vias_list = []
    id_via = 1

    for r in range(9, 2944):
        cod = clean_str(ws.cell(row=r, column=1).value)
        tipo = clean_str(ws.cell(row=r, column=2).value)
        nom = clean_str(ws.cell(row=r, column=3).value)
        obs = clean_str(ws.cell(row=r, column=4).value)
        sec = clean_str(ws.cell(row=r, column=5).value)
        cond = clean_str(ws.cell(row=r, column=6).value)
        clasif = clean_str(ws.cell(row=r, column=7).value)

        if not cod and not nom:
            continue
        if cod.startswith("TOTAL") or cod.startswith("TIPO") or nom.startswith("TOTAL"):
            continue

        nom_clean = re.sub(r"\s+", " ", nom).strip().upper()
        if not nom_clean:
            continue

        id_tipo = tipo_via_map.get(tipo, 2)
        variants = set()
        variants.add(nom_clean)
        variants.add(remove_accents(nom_clean))

        # Base sin prefijos redundantes pegados en el nombre (ej. CA. PACASMAYO -> PACASMAYO)
        nom_base = re.sub(
            r"^(?:CA\.|AV\.|JR\.|PJE\.|PRLG\.|CTRA\.|CALLE|AVENIDA|JIRON|PASAJE|PROLONGACION)\s+",
            "",
            nom_clean,
            flags=re.IGNORECASE,
        ).strip()
        if nom_base and nom_base != nom_clean:
            variants.add(nom_base)
            variants.add(remove_accents(nom_base))

        # Variantes de paréntesis ej: PACIFICO (JUAN TOMIS STACK)
        for target_nom in [nom_clean, nom_base]:
            m = re.match(r"^(.*?)\s*\((.*?)\)$", target_nom)
            if m:
                p1 = m.group(1).strip()
                p2 = m.group(2).strip()
                variants.update([
                    p1,
                    remove_accents(p1),
                    p2,
                    remove_accents(p2),
                    f"{p1} {p2}",
                    remove_accents(f"{p1} {p2}"),
                ])

        # Variantes de guiones ej: CHICLAYO - FERREÑAFE
        for target_nom in [nom_clean, nom_base]:
            if "-" in target_nom:
                sin_guion = re.sub(r"\s*-\s*", " ", target_nom)
                variants.update([sin_guion, remove_accents(sin_guion)])
                partes = [p.strip() for p in re.split(r"\s*-\s*", target_nom) if p.strip()]
                if len(partes) == 2:
                    variants.update([
                        f"{partes[0]} A {partes[1]}",
                        f"{partes[1]} A {partes[0]}",
                        partes[1],
                        remove_accents(partes[1]),
                    ])

        # Observación o denominación anterior / alterna
        if obs:
            obs_clean = re.sub(r"\s+", " ", obs).strip().upper()
            variants.update([obs_clean, remove_accents(obs_clean)])

        # Variantes de números y fechas (ej. 7 DE ENERO <-> SIETE DE ENERO)
        for target_nom in list(variants):
            for num_k, num_repls in number_words_map.items():
                if num_k in target_nom:
                    for rep in num_repls:
                        variants.update([
                            target_nom.replace(num_k, rep),
                            remove_accents(target_nom.replace(num_k, rep)),
                        ])

        # Variantes con nombres de personajes / abreviaturas
        for target_nom in [nom_clean, nom_base]:
            words = target_nom.split()
            if len(words) >= 3:
                rest = " ".join(words[2:])
                variants.update([
                    f"{words[0]} {words[1][0]}. {rest}",
                    f"{words[0]} {words[1][0]} {rest}",
                    f"{words[0]} {words[-1]}",
                ])
                if len(words[-1]) >= 5 and words[-1] not in ("NORTE", "SUR", "ESTE", "OESTE"):
                    variants.add(words[-1])
            elif len(words) == 2:
                variants.update([
                    f"{words[0][0]}. {words[1]}",
                    f"{words[0][0]} {words[1]}",
                ])
                if len(words[1]) >= 5 and words[0] not in ("SIN", "SAN", "DE", "DEL", "LA", "EL", "LOS", "LAS", "SUR", "NORTE"):
                    variants.add(words[1])

        # Santos / Religiosos
        for item in list(variants):
            if "SAN " in item:
                variants.update([item.replace("SAN ", "SN. "), item.replace("SAN ", "SN ")])
            if "SANTA " in item:
                variants.update([item.replace("SANTA ", "STA. "), item.replace("SANTA ", "STA ")])
            if "SANTO " in item:
                variants.update([item.replace("SANTO ", "STO. "), item.replace("SANTO ", "STO ")])
            if "SEÑOR DE " in item:
                variants.update([item.replace("SEÑOR DE ", "SR. DE "), item.replace("SEÑOR DE ", "SR DE ")])

        # Variantes con prefijos de tipo de vía
        prefixes = via_prefixes.get(tipo, [tipo])
        cands = list(set([nom_clean, remove_accents(nom_clean), nom_base, remove_accents(nom_base)]))
        for c in list(cands):
            for num_k, num_repls in number_words_map.items():
                if num_k in c:
                    for rep in num_repls:
                        cands.append(c.replace(num_k, rep))

        for cand in cands:
            for pfx in prefixes:
                variants.update([f"{pfx} {cand}", remove_accents(f"{pfx} {cand}")])

        clean_variants = []
        for v in variants:
            v_s = re.sub(r"\s+", " ", v).strip().upper()
            v_s = re.sub(r"[\(\)]", "", v_s).strip()
            if len(v_s) > 1 and v_s != nom_clean and v_s not in clean_variants:
                clean_variants.append(v_s)

        synonyms = [nom_clean] + sorted(clean_variants)

        vias_list.append({
            "id": id_via,
            "codigo": cod,
            "id_tipo_via": id_tipo,
            "nom_via": nom_clean,
            "sinonimos": synonyms,
            "sector": sec,
            "condicion": cond,
            "clasificacion": clasif,
            "obs_alterna": obs if obs else None,
        })
        id_via += 1

    logger.info("Total vías oficiales extraídas: %d", len(vias_list))
    return vias_list


def extract_zonas_habilitaciones_urbanas():
    """Extrae las 460 zonas oficiales de la hoja 2 FORMATO CODIFICADOR HU."""
    tipo_zona_map = {
        "URB.": 6,
        "P.J.": 5,
        "CONJ.RES.": 4,
        "CONJ.HAB.": 3,
        "A.H.": 1,
        "ASOC.VIV.": 21,
        "ASOC.PVIV.": 24,
        "COOP.VIV.": 22,
        "LOTIZ.": 12,
        "UPIS": 7,
        "H.U.": 6,
        "URB.PROG.": 6,
        "ASOC.PROP.": 10,
        "ASOC.POBL.": 10,
        "HAB.IND.": 6,
        "P.I.": 6,
        "P.A.L.": 1,
        "P.T.": 6,
        "C.P.M.": 26,
        "FDO.": 19,
    }

    zona_prefixes = {
        "URB.": ["URB.", "URB", "URBANIZACION", "URBANIZACIÓN"],
        "P.J.": ["P.J.", "PJ", "P.J", "PUEBLO JOVEN", "PP.JJ."],
        "A.H.": ["A.H.", "AH", "A.H", "AA.HH.", "AAHH", "ASENTAMIENTO HUMANO"],
        "CONJ.RES.": ["CONJ.RES.", "CONJ RES", "CONJUNTO RESIDENCIAL", "RESIDENCIAL"],
        "CONJ.HAB.": ["CONJ.HAB.", "CONJ HAB", "CONJUNTO HABITACIONAL"],
        "ASOC.VIV.": ["ASOC.VIV.", "ASOC VIV", "ASOCIACION DE VIVIENDA", "ASOCIACIÓN DE VIVIENDA"],
        "ASOC.PVIV.": ["ASOC.PVIV.", "ASOC PVIV", "ASOC. PRO VIVIENDA", "ASOCIACION PRO VIVIENDA"],
        "COOP.VIV.": ["COOP.VIV.", "COOP VIV", "COOPERATIVA DE VIVIENDA"],
        "LOTIZ.": ["LOTIZ.", "LOTIZ", "LOTIZACION", "LOTIZACIÓN"],
        "UPIS": ["UPIS", "U.P.I.S."],
        "H.U.": ["H.U.", "HU", "HABILITACION URBANA", "HABILITACIÓN URBANA", "URB.", "URB"],
        "URB.PROG.": ["URB.PROG.", "URB. PROG.", "URBANIZACION POPULAR", "URB."],
        "FDO.": ["FDO.", "FDO", "FUNDO"],
    }

    logger.info("Cargando libro de zonas: %s", EXCEL_ZONAS.name)
    wb = openpyxl.load_workbook(str(EXCEL_ZONAS), data_only=True)
    ws = wb["2 FORMATO CODIFICADOR HU"]

    zonas_list = []
    id_zona = 1

    for r in range(8, ws.max_row + 1):
        cod = clean_str(ws.cell(row=r, column=1).value)
        tipo = clean_str(ws.cell(row=r, column=2).value)
        nom = clean_str(ws.cell(row=r, column=3).value)
        sec = clean_str(ws.cell(row=r, column=4).value)
        cond = clean_str(ws.cell(row=r, column=5).value)

        if not cod:
            continue
        if cod.startswith("TIPOS") or cod.startswith("TOTAL"):
            continue

        nom_clean = re.sub(r"\s+", " ", nom).strip().upper()
        if not nom_clean:
            continue

        id_tipo = tipo_zona_map.get(tipo, 6)
        variants = set()
        variants.add(nom_clean)
        variants.add(remove_accents(nom_clean))

        # Guiones y etapas
        if "-" in nom_clean:
            sin_guion = re.sub(r"\s*-\s*", " ", nom_clean)
            variants.add(sin_guion)
            variants.add(remove_accents(sin_guion))
            partes = [p.strip() for p in re.split(r"\s*-\s*", nom_clean) if p.strip()]
            if len(partes) >= 2:
                variants.add(partes[0])
                variants.add(remove_accents(partes[0]))
                if "ETAPA" in partes[-1] or re.match(r"^(?:I|II|III|IV|V)\b", partes[-1]):
                    variants.add(f"{partes[0]} {partes[-1]}")
                    variants.add(remove_accents(f"{partes[0]} {partes[-1]}"))

        # Conjugaciones y variantes de etapas
        etapa_map = {
            "I ETAPA": ["1 ETAPA", "ETAPA 1", "ETAPA I", "1RA ETAPA", "PRIMERA ETAPA"],
            "II ETAPA": ["2 ETAPA", "ETAPA 2", "ETAPA II", "2DA ETAPA", "SEGUNDA ETAPA"],
            "III ETAPA": ["3 ETAPA", "ETAPA 3", "ETAPA III", "3RA ETAPA", "TERCERA ETAPA"],
            "IV ETAPA": ["4 ETAPA", "ETAPA 4", "ETAPA IV", "4TA ETAPA", "CUARTA ETAPA"],
        }
        for item in list(variants):
            for rom, replacements in etapa_map.items():
                if rom in item:
                    for rep in replacements:
                        variants.add(item.replace(rom, rep))
                        variants.add(remove_accents(item.replace(rom, rep)))

        # Santos / Religiosos
        for item in list(variants):
            if "SANTA " in item:
                variants.add(item.replace("SANTA ", "STA. "))
                variants.add(item.replace("SANTA ", "STA "))
            if "SAN " in item:
                variants.add(item.replace("SAN ", "SN. "))
                variants.add(item.replace("SAN ", "SN "))
            if "SEÑOR DE " in item:
                variants.add(item.replace("SEÑOR DE ", "SR. DE "))
                variants.add(item.replace("SEÑOR DE ", "SR DE "))

        # Prefijos de tipo de zona
        prefixes = zona_prefixes.get(tipo, [tipo])
        base_cands = [nom_clean, remove_accents(nom_clean)]
        if " - " in nom_clean:
            base_cands.append(nom_clean.split(" - ")[0].strip())
            base_cands.append(remove_accents(nom_clean.split(" - ")[0].strip()))

        for cand in base_cands:
            for pfx in prefixes:
                variants.add(f"{pfx} {cand}")
                variants.add(remove_accents(f"{pfx} {cand}"))

        clean_variants = []
        for v in variants:
            v_s = re.sub(r"\s+", " ", v).strip().upper()
            v_s = re.sub(r"[\(\)]", "", v_s).strip()
            if len(v_s) > 1 and v_s != nom_clean and v_s not in clean_variants:
                clean_variants.append(v_s)

        synonyms = [nom_clean] + sorted(clean_variants)

        zonas_list.append({
            "id": id_zona,
            "codigo": cod,
            "id_tipo_zona": id_tipo,
            "nom_zona": nom_clean,
            "sinonimos": synonyms,
            "sector": sec,
            "condicion": cond,
        })
        id_zona += 1

    logger.info("Total zonas oficiales extraídas: %d", len(zonas_list))
    return zonas_list


def main():
    logger.info("Iniciando extracción y normalización de catálogos oficiales...")

    vias = extract_vias_oficiales()
    with open(OUTPUT_VIAS_JSON, "w", encoding="utf-8") as f:
        json.dump(vias, f, ensure_ascii=False, indent=2)
    logger.info("Guardado exitosamente: %s con %d vías", OUTPUT_VIAS_JSON.name, len(vias))

    zonas = extract_zonas_habilitaciones_urbanas()
    with open(OUTPUT_ZONAS_JSON, "w", encoding="utf-8") as f:
        json.dump(zonas, f, ensure_ascii=False, indent=2)
    logger.info("Guardado exitosamente: %s con %d zonas", OUTPUT_ZONAS_JSON.name, len(zonas))


if __name__ == "__main__":
    main()
