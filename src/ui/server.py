"""Servidor backend FastAPI con soporte SSE (Server-Sent Events) para la GUI del ETL MPCH."""

import asyncio
import json
import logging
import os
import queue
import signal
import threading
import time
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException, Query, Request, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from src.config.settings import DatabaseSettings, OllamaSettings, get_settings
from src.extractors.db_extractor import DatabaseExtractor
from src.pipelines.etl_pipeline import ETLPipeline
from src.services.db_service import DatabaseService
from src.services.ollama_service import OllamaService
from src.transformers.ai_parser import AIAddressParser
from src.transformers.pipeline_transformer import PipelineTransformer

logger = logging.getLogger("etl_mpch.gui")

# Configuración del servidor y directorios estáticos
STATIC_DIR = Path(__file__).resolve().parent / "static"


class ExecutionState:
    """Estado compartido en memoria de la ejecución del pipeline."""

    def __init__(self):
        settings = get_settings()
        self.is_running: bool = False
        self.pipeline: Optional[ETLPipeline] = None
        self.thread: Optional[threading.Thread] = None
        self.active_schema: str = settings.db.schema
        self.active_table: str = settings.db.table
        self.num_workers: int = 4
        self.active_model: Optional[str] = None
        # Conexiones dinámicas establecidas por el wizard (independientes de .env)
        self.dynamic_db_service: Optional[DatabaseService] = None
        self.dynamic_db_settings: Optional[Any] = None
        self.dynamic_ollama_service: Optional[OllamaService] = None
        self.dynamic_ollama_settings: Optional[Any] = None
        self.recent_records: List[Dict[str, Any]] = []
        self.max_recent_records: int = 500
        self.log_history: List[str] = []
        self.max_log_history: int = 1000
        self.stats: Dict[str, Any] = {
            "status": "IDLE",
            "total_to_process": 0,
            "processed": 0,
            "valid": 0,
            "observed": 0,
            "failed": 0,
            "ai_records": 0,
            "hybrid_records": 0,
            "heuristic_records": 0,
            "progress_pct": 0.0,
            "start_time": None,
            "elapsed_seconds": 0,
        }
        self.clients: List[asyncio.Queue] = []
        self._lock = threading.Lock()
        self.loop: Optional[asyncio.AbstractEventLoop] = None

    def reset_for_run(
        self,
        total: int,
        schema: str,
        table: str,
        num_workers: int = 4,
        model: Optional[str] = None,
    ):
        with self._lock:
            self.is_running = True
            self.active_schema = schema
            self.active_table = table
            self.num_workers = max(1, int(num_workers or 1))
            self.active_model = model
            # No vaciar self.recent_records para preservar el historial acumulado de la sesión
            self.stats = {
                "status": "RUNNING",
                "total_to_process": total,
                "processed": 0,
                "valid": 0,
                "observed": 0,
                "failed": 0,
                "ai_records": 0,
                "hybrid_records": 0,
                "heuristic_records": 0,
                "progress_pct": 0.0,
                "start_time": time.time(),
                "elapsed_seconds": 0,
            }

    def clear_session(self):
        with self._lock:
            self.recent_records.clear()
            self.log_history.clear()
            self.stats = {
                "status": "IDLE",
                "total_to_process": 0,
                "processed": 0,
                "valid": 0,
                "observed": 0,
                "failed": 0,
                "ai_records": 0,
                "hybrid_records": 0,
                "heuristic_records": 0,
                "progress_pct": 0.0,
                "start_time": None,
                "elapsed_seconds": 0,
            }

    def add_log(self, message: str):
        with self._lock:
            self.log_history.append(message)
            if len(self.log_history) > self.max_log_history:
                self.log_history.pop(0)
        self.broadcast({"type": "log", "payload": message})

    def add_record(self, record_data: Dict[str, Any]):
        idx = record_data.get("index", 0)
        tot = record_data.get("total", 0)
        pk = record_data.get("id_licencia", "")
        raw = record_data.get("raw_text", "")
        motor = record_data.get("metodo", "IA")
        es_proc = record_data.get("es_procesado", False)
        obs = record_data.get("observacion", "")
        via = record_data.get("via_desc") or f"{record_data.get('nom_via') or 'N/D'} N° {record_data.get('num_via') or 'S/N'}"
        zona = record_data.get("zona_desc") or (record_data.get("nom_zona") or "N/D")
        cat = record_data.get("catastro") or ""
        ref = record_data.get("referencia") or ""
        status_tag = "VALIDO" if es_proc else "OBSERVADO"

        log_line = f"[{idx}/{tot} | ID: {pk}] [{status_tag}] [{motor}] \"{raw}\" -> {via} | {zona}"
        if cat:
            log_line += f" | {cat}"
        if ref:
            log_line += f" | Ref: {ref}"
        if obs:
            log_line += f" | Motivo: {obs}"

        with self._lock:
            # Actualizar in-place si el registro ya existía en la sesión
            updated_existing = False
            for i, existing in enumerate(self.recent_records):
                if existing.get("id_licencia") == pk:
                    self.recent_records[i] = record_data
                    updated_existing = True
                    break

            if not updated_existing:
                self.recent_records.append(record_data)
                if len(self.recent_records) > self.max_recent_records:
                    self.recent_records.pop(0)

            self.log_history.append(log_line)
            if len(self.log_history) > self.max_log_history:
                self.log_history.pop(0)

            total = self.stats.get("total_to_process", 0)
            processed = record_data.get("processed_count", 0)
            pct = round((processed / total * 100), 1) if total > 0 else 0.0

            start_t = self.stats.get("start_time")
            elapsed = round(time.time() - start_t, 1) if start_t else 0
            speed_rps = round(processed / elapsed, 2) if elapsed > 0 else 0.0

            self.stats.update({
                "processed": processed,
                "valid": record_data.get("valid_count", 0),
                "observed": record_data.get("observed_count", 0),
                "failed": record_data.get("failed_count", 0),
                "ai_records": record_data.get("ai_records", 0),
                "hybrid_records": record_data.get("hybrid_records", 0),
                "heuristic_records": record_data.get("heuristic_records", 0),
                "progress_pct": pct,
                "elapsed_seconds": elapsed,
                "speed_rps": speed_rps,
                "num_workers": self.num_workers,
            })

        self.broadcast({"type": "record", "payload": record_data})
        self.broadcast({"type": "stats", "payload": self.stats})
        self.broadcast({"type": "log", "payload": log_line})

    def finish_run(self, summary_data: Optional[Dict[str, Any]] = None, error: Optional[str] = None):
        with self._lock:
            self.is_running = False
            self.pipeline = None
            self.stats["status"] = "ERROR" if error else "FINISHED"
            if error:
                self.stats["error"] = error
        if error:
            self.broadcast({"type": "error", "payload": {"error": error}})
        else:
            self.broadcast({"type": "done", "payload": summary_data or self.stats})

    def register_client(self, q: asyncio.Queue):
        with self._lock:
            self.clients.append(q)

    def unregister_client(self, q: asyncio.Queue):
        with self._lock:
            if q in self.clients:
                self.clients.remove(q)

    def broadcast(self, event: Dict[str, Any]):
        if not self.loop:
            return
        with self._lock:
            targets = list(self.clients)
        for q in targets:
            try:
                self.loop.call_soon_threadsafe(q.put_nowait, event)
            except Exception:
                pass

    def shutdown(self):
        """Detiene pipelines activos y despierta/cierra todas las colas de clientes SSE."""
        with self._lock:
            if self.is_running and self.pipeline:
                try:
                    self.pipeline.request_stop()
                except Exception:
                    pass
            targets = list(self.clients)
            self.clients.clear()

        for q in targets:
            try:
                if self.loop and self.loop.is_running():
                    self.loop.call_soon_threadsafe(q.put_nowait, None)
                else:
                    q.put_nowait(None)
            except Exception:
                pass


state = ExecutionState()


@asynccontextmanager
async def lifespan(app_instance: FastAPI):
    """Ciclo de vida de FastAPI: registra el event loop activo en el estado."""
    state.loop = asyncio.get_running_loop()
    try:
        yield
    finally:
        state.shutdown()


app = FastAPI(
    title="MPCH ETL - Interfaz Gráfica de Normalización",
    description="Panel interactivo de control y visualización registro por registro con IA",
    version="2.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def add_no_cache_headers(request: Request, call_next):
    response: Response = await call_next(request)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


# ── Modelos para el flujo wizard ───────────────────────────────────────────────

class AIConnectRequest(BaseModel):
    host: str = Field(default="localhost", description="Host o IP del servidor Ollama")
    port: int = Field(default=11434, description="Puerto de conexión a Ollama")
    model: str = Field(default="patroclo-artesano-7b:latest", description="Nombre del modelo seleccionado")
    timeout: Optional[int] = Field(default=60, description="Timeout en segundos")
    temperature: Optional[float] = Field(default=0.0, description="Temperatura de inferencia")
    num_workers: Optional[int] = Field(default=4, ge=1, le=16, description="Número de trabajadores concurrentes a Ollama")


class DetectModelsRequest(BaseModel):
    host: str = Field(default="localhost", description="Host o IP del servidor Ollama")
    port: int = Field(default=11434, description="Puerto de conexión a Ollama")


class ConnectRequest(BaseModel):
    host: str = Field(default="localhost", description="Host del servidor PostgreSQL")
    port: int = Field(default=5432, description="Puerto de conexión")
    dbname: str = Field(default="bd_mpch", description="Nombre de la base de datos")
    user: str = Field(default="postgres", description="Usuario de PostgreSQL")
    password: str = Field(default="postgres", description="Contraseña del usuario")


class InspectSchemaRequest(BaseModel):
    schema_name: str = Field(..., description="Nombre del esquema a inspeccionar")


class InspectTableRequest(BaseModel):
    schema_name: str = Field(..., description="Nombre del esquema")
    table_name: str = Field(..., description="Nombre de la tabla de direcciones")
    id_col: Optional[str] = Field(default=None, description="Columna que actúa como ID/PK")
    address_col: Optional[str] = Field(default=None, description="Columna con la dirección completa")


class StartPipelineRequest(BaseModel):
    schema_name: str = Field(default_factory=lambda: get_settings().db.schema, description="Esquema destino en PostgreSQL")
    table_name: str = Field(default_factory=lambda: get_settings().db.table, description="Nombre de tabla de direcciones")
    id_col: Optional[str] = Field(default=None, description="Columna ID a conservar")
    address_col: Optional[str] = Field(default=None, description="Columna de dirección a normalizar")
    limit: Optional[int] = Field(default=None, description="Límite máximo de registros a procesar")
    batch_size: Optional[int] = Field(default=50, description="Tamaño de lote")
    filter_mode: Optional[str] = Field(default="pending", description="Modo de filtro: pending, all, observed")
    require_ai: bool = Field(default=True, description="Si es True, falla de inmediato si Ollama no está operativo")
    num_workers: Optional[int] = Field(default=4, ge=1, le=16, description="Número de peticiones concurrentes a Ollama")


@app.get("/", response_class=FileResponse)
async def serve_ui():
    """Sirve la página de interfaz gráfica de escritorio sin caché."""
    index_path = STATIC_DIR / "index.html"
    if not index_path.exists():
        raise HTTPException(status_code=404, detail="Archivo UI index.html no encontrado.")
    return FileResponse(
        str(index_path),
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )


@app.get("/api/status")
async def get_system_status():
    """Retorna el estado del sistema, conexiones, esquemas y métricas actuales."""
    settings = get_settings()
    db_svc = state.dynamic_db_service or DatabaseService()
    ollama_svc = state.dynamic_ollama_service or OllamaService()

    db_health = db_svc.check_connection(schema=state.active_schema)
    ai_health = ollama_svc.check_connection()

    table_counts = {}
    try:
        extractor = DatabaseExtractor(
            db_service=db_svc,
            schema=state.active_schema,
            table=state.active_table,
        )
        table_counts = extractor.get_status_counts()
    except Exception as e:
        logger.warning("No se pudo obtener conteo de tabla: %s", e)
        table_counts = {"total": 0, "pendientes": 0, "validos": 0, "observados": 0}

    current_model = state.active_model or (ai_health.get("target_model") or ai_health.get("model") or ollama_svc.model_name)

    return {
        "running": state.is_running,
        "is_running": state.is_running,
        "active_schema": state.active_schema,
        "active_table": state.active_table,
        "active_model": current_model,
        "num_workers": state.num_workers,
        "has_completed_session": (state.stats.get("status") == "FINISHED" and len(state.recent_records) > 0),
        "available_schemas": db_health.get("available_schemas", ["public"]),
        "db": {
            "connected": db_health.get("connected", False),
            "schema": db_health.get("schema", state.active_schema),
            "table_exists": db_health.get("table_exists", False),
            "message": db_health.get("message", ""),
        },
        "ai": {
            "connected": ai_health.get("connected", False),
            "model": current_model,
            "installed": ai_health.get("model_available", False) or ai_health.get("model_installed", False),
            "model_available": ai_health.get("model_available", False) or ai_health.get("model_installed", False),
            "base_url": ollama_svc.base_url,
            "message": ai_health.get("message", ""),
        },
        "table_counts": table_counts,
        "stats": state.stats,
        "recent_records": state.recent_records[-100:],
        "log_history": state.log_history[-200:],
    }


@app.get("/api/schemas")
async def get_schemas():
    """Devuelve la lista de esquemas accesibles en la base de datos."""
    try:
        db_svc = DatabaseService()
        schemas = db_svc.get_available_schemas()
        return {"schemas": schemas}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error consultando esquemas: {e}")


@app.get("/api/counts")
async def get_table_counts(
    schema: str = Query("public"),
    table: str = Query("direcciones_actual")
):
    """Devuelve los conteos catastrales de la tabla especificada."""
    try:
        db_svc = state.dynamic_db_service or DatabaseService()
        extractor = DatabaseExtractor(db_service=db_svc, schema=schema, table=table)
        counts = extractor.get_status_counts()
        return counts
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error obteniendo conteo: {e}")


# ── Endpoints del Wizard ────────────────────────────────────────────────────────

@app.post("/api/detect-models")
async def wizard_detect_models(req: DetectModelsRequest):
    """Paso 1 del wizard: consulta el servidor Ollama y lista los modelos descargados."""
    import httpx
    base_url = f"http://{req.host}:{req.port}".rstrip("/")
    try:
        async with httpx.AsyncClient(timeout=4.0) as client:
            resp = await client.get(f"{base_url}/api/tags")
        if resp.status_code == 200:
            data = resp.json()
            models = [m.get("name") for m in data.get("models", [])]
            return {
                "connected": True,
                "host": req.host,
                "port": req.port,
                "models": models,
                "count": len(models),
                "message": f"Conexión exitosa. {len(models)} modelo(s) encontrado(s).",
            }
        else:
            return {
                "connected": False,
                "host": req.host,
                "port": req.port,
                "models": [],
                "count": 0,
                "message": f"Servidor Ollama respondió con código HTTP {resp.status_code}",
            }
    except Exception as e:
        return {
            "connected": False,
            "host": req.host,
            "port": req.port,
            "models": [],
            "count": 0,
            "message": f"No se pudo conectar a Ollama en {base_url}: {str(e)}",
        }


@app.post("/api/connect-ai")
async def wizard_connect_ai(req: AIConnectRequest):
    """Paso 1 del wizard: prueba y guarda la conexión al modelo de IA con parámetros dinámicos."""
    import time
    base_url = f"http://{req.host}:{req.port}".rstrip("/")
    dyn_settings = OllamaSettings(
        OLLAMA_BASE_URL=base_url,
        OLLAMA_MODEL=req.model,
        OLLAMA_TIMEOUT=req.timeout or 60,
        OLLAMA_TEMPERATURE=req.temperature or 0.0,
    )
    ollama_svc = OllamaService(settings=dyn_settings)

    t0 = time.perf_counter()
    health = ollama_svc.check_connection()
    latency_ms = round((time.perf_counter() - t0) * 1000, 1)

    if not health.get("connected"):
        raise HTTPException(
            status_code=400,
            detail=f"No se pudo conectar al servidor Ollama en {base_url}. Verifique que Ollama esté ejecutándose."
        )

    available_models = health.get("available_models", [])
    model_match = health.get("model_available", False)

    if not model_match:
        model_list_str = ", ".join(available_models) if available_models else "ninguno"
        raise HTTPException(
            status_code=400,
            detail=(
                f"Conexión establecida con Ollama en {base_url}, pero el modelo '{req.model}' "
                f"no está disponible en el servidor. Modelos detectados: [{model_list_str}]."
            )
        )

    # Guardar en estado compartido de la sesión
    state.dynamic_ollama_service = ollama_svc
    state.dynamic_ollama_settings = dyn_settings
    if req.num_workers:
        state.num_workers = max(1, min(16, int(req.num_workers)))

    return {
        "success": True,
        "host": req.host,
        "port": req.port,
        "model": req.model,
        "num_workers": state.num_workers,
        "model_available": True,
        "available_models": available_models,
        "latency_ms": latency_ms,
        "message": f"Conexión exitosa con Ollama ({req.model}) en {base_url} ({latency_ms} ms)",
    }


@app.post("/api/connect")
async def wizard_connect(req: ConnectRequest):
    """Paso 2 del wizard: prueba la conexión con parámetros dinámicos y retorna esquemas disponibles."""
    from src.config.settings import DatabaseSettings

    try:
        dyn_settings = DatabaseSettings(
            DB_HOST=req.host,
            DB_PORT=req.port,
            DB_NAME=req.dbname,
            DB_USER=req.user,
            DB_PASSWORD=req.password,
        )
        db_svc = DatabaseService(settings=dyn_settings)
        schemas = db_svc.get_available_schemas()

        if not schemas:
            # Si get_available_schemas retorna vacío, intentar verificar conexión
            health = db_svc.check_connection()
            if not health.get("connected"):
                raise HTTPException(status_code=400, detail=health.get("message", "No se pudo conectar a la base de datos."))
            schemas = health.get("available_schemas", [])

        # Guardar la conexión en el estado compartido
        state.dynamic_db_service = db_svc
        state.dynamic_db_settings = dyn_settings

        return {
            "success": True,
            "message": f"Conexión exitosa a '{req.dbname}' en {req.host}:{req.port}",
            "schemas": schemas,
            "host": req.host,
            "port": req.port,
            "dbname": req.dbname,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error de conexión: {str(e)}")


@app.post("/api/inspect-schema")
async def wizard_inspect_schema(req: InspectSchemaRequest):
    """Paso 2 del wizard: lista tablas del esquema seleccionado con columnas y cantidad de filas."""
    db_svc = state.dynamic_db_service or DatabaseService()

    try:
        from sqlalchemy import text as sa_text

        tables_info = []
        with db_svc.get_session() as session:
            # Consultar tablas del esquema
            rows = session.execute(sa_text("""
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = :schema
                  AND table_type = 'BASE TABLE'
                ORDER BY table_name;
            """), {"schema": req.schema_name}).fetchall()

            for (tname,) in rows:
                # Obtener columnas de cada tabla
                col_rows = session.execute(sa_text("""
                    SELECT column_name, data_type
                    FROM information_schema.columns
                    WHERE table_schema = :schema AND table_name = :table
                    ORDER BY ordinal_position;
                """), {"schema": req.schema_name, "table": tname}).fetchall()

                columns = [{"name": c[0], "type": c[1]} for c in col_rows]

                # Obtener conteo de filas
                try:
                    count_val = session.execute(
                        sa_text(f'SELECT COUNT(*) FROM "{req.schema_name}"."{tname}";')
                    ).scalar() or 0
                except Exception:
                    count_val = -1

                tables_info.append({
                    "name": tname,
                    "columns": columns,
                    "row_count": count_val,
                })

        return {
            "schema": req.schema_name,
            "tables": tables_info,
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error inspeccionando esquema '{req.schema_name}': {str(e)}")


@app.post("/api/inspect-table")
async def wizard_inspect_table(req: InspectTableRequest):
    """Paso 3 del wizard: preview de filas y conteos de estado de la tabla de direcciones seleccionada."""
    db_svc = state.dynamic_db_service or DatabaseService()
    settings = get_settings()

    id_col = req.id_col or settings.db.id_col
    address_col = req.address_col or settings.db.dir_col

    try:
        from sqlalchemy import text as sa_text

        with db_svc.get_session() as session:
            # Obtener preview de primeras 10 filas con las columnas clave
            preview_rows = session.execute(sa_text(f"""
                SELECT "{id_col}", "{address_col}"
                FROM "{req.schema_name}"."{req.table_name}"
                LIMIT 10;
            """)).fetchall()

            preview = [{"id": r[0], "direccion": r[1]} for r in preview_rows]

            # Conteo total
            total = session.execute(
                sa_text(f'SELECT COUNT(*) FROM "{req.schema_name}"."{req.table_name}";')
            ).scalar() or 0

            # Intentar obtener conteos de estado si la columna es_procesado existe
            try:
                counts = session.execute(sa_text(f"""
                    SELECT
                        COUNT(*) FILTER (WHERE es_procesado IS NULL)     AS pendientes,
                        COUNT(*) FILTER (WHERE es_procesado = TRUE)       AS validos,
                        COUNT(*) FILTER (WHERE es_procesado = FALSE)      AS observados
                    FROM "{req.schema_name}"."{req.table_name}";
                """)).fetchone()
                status_counts = {
                    "total": total,
                    "pendientes": counts[0] if counts else total,
                    "validos": counts[1] if counts else 0,
                    "observados": counts[2] if counts else 0,
                }
            except Exception:
                status_counts = {"total": total, "pendientes": total, "validos": 0, "observados": 0}

        return {
            "schema": req.schema_name,
            "table": req.table_name,
            "id_col": id_col,
            "address_col": address_col,
            "counts": status_counts,
            "preview": preview,
        }
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Error al inspeccionar tabla '{req.schema_name}.{req.table_name}': {str(e)}"
        )


@app.post("/api/test-ai")
async def test_ai_connection(sample_address: Optional[str] = None):
    """Ejecuta una prueba de inferencia en vivo con el modelo Ollama configurado."""
    ollama_svc = state.dynamic_ollama_service or OllamaService()
    test_addr = sample_address or "CALLE SAN JOSE 456 URB SANTA VICTORIA CHICLAYO"

    start_time = time.perf_counter()
    health = ollama_svc.check_connection()
    if not health.get("connected"):
        return {
            "success": False,
            "model": ollama_svc.model_name,
            "latency_ms": 0,
            "message": health.get("message", "Ollama fuera de línea"),
        }

    try:
        result = ollama_svc.parse_address_with_ai(test_addr)
        elapsed_ms = round((time.perf_counter() - start_time) * 1000, 1)
        parsed_dict = result.model_dump() if hasattr(result, "model_dump") else (result.dict() if result else None)
        return {
            "success": bool(result),
            "model": ollama_svc.model_name,
            "latency_ms": elapsed_ms,
            "input": test_addr,
            "parsed": parsed_dict,
            "message": f"Inferencia exitosa en {elapsed_ms} ms" if result else "Modelo respondió pero no extrajo estructura",
        }
    except Exception as e:
        return {
            "success": False,
            "model": ollama_svc.model_name,
            "latency_ms": 0,
            "message": f"Error de inferencia: {e}",
        }


def _run_pipeline_worker(req: StartPipelineRequest):
    """Función que ejecuta el pipeline ETL en un hilo secundario para no bloquear el servidor."""
    settings = get_settings()
    target_schema = req.schema_name or settings.db.schema
    target_table  = req.table_name  or settings.db.table
    state.add_log(f"Iniciando pipeline ETL en esquema [{target_schema}] tabla [{target_table}]...")

    try:
        # Usar conexiones dinámicas del wizard si están disponibles
        db_svc = state.dynamic_db_service or DatabaseService()
        ollama_svc = state.dynamic_ollama_service or OllamaService()

        # Resolver columnas: las del wizard tienen prioridad sobre las del .env
        id_col   = req.id_col      or settings.db.id_col
        dir_col  = req.address_col or settings.db.dir_col

        state.add_log(f"Columna ID: [{id_col}] | Columna dirección: [{dir_col}]")
        state.add_log(f"Motor IA: [{ollama_svc.base_url}] | Modelo: [{ollama_svc.model_name}]")

        # Construir el extractor con las columnas identificadas por el usuario en el wizard
        custom_extractor = DatabaseExtractor(
            db_service=db_svc,
            schema=target_schema,
            table=target_table,
            id_col=id_col,
            dir_col=dir_col,
        )

        # Construir el transformador con el servicio de IA dinámico
        custom_ai_parser = AIAddressParser(ollama_service=ollama_svc)
        custom_transformer = PipelineTransformer(
            ai_parser=custom_ai_parser,
            use_ai=req.require_ai,
        )

        workers_count = max(1, min(16, int(req.num_workers or state.num_workers or 4)))
        state.num_workers = workers_count
        state.active_model = ollama_svc.model_name

        pipeline = ETLPipeline(
            schema=target_schema,
            table=target_table,
            batch_size=req.batch_size,
            db_service=db_svc,
            require_ai=req.require_ai,
            extractor=custom_extractor,
            transformer=custom_transformer,
            on_record_processed=state.add_record,
            num_workers=workers_count,
        )
        state.pipeline = pipeline

        state.add_log("Preparando entorno y tablas maestras...")
        pipeline.prepare_environment()

        # Calcular número total a procesar
        extractor = pipeline.extractor
        status_counts = extractor.get_status_counts()
        total_avail = status_counts.get("total", 0)
        pending_avail = status_counts.get("pendientes", 0)
        observed_avail = status_counts.get("observados", 0)

        # Normalizar modo de filtro (soporta 'pending', 'observed', 'all', 'unprocessed')
        raw_filter = str(req.filter_mode or "pending").strip().lower()
        if raw_filter in ("unprocessed", "observados", "observed"):
            filter_mode = "observed"
        elif raw_filter in ("pending", "pendientes"):
            filter_mode = "pending"
        else:
            filter_mode = "all"

        if filter_mode == "pending":
            target_pool = pending_avail
        elif filter_mode == "observed":
            target_pool = observed_avail
        else:
            target_pool = total_avail

        to_process = min(req.limit, target_pool) if req.limit else target_pool
        state.reset_for_run(
            total=to_process,
            schema=target_schema,
            table=target_table,
            num_workers=workers_count,
            model=ollama_svc.model_name,
        )
        state.add_log(f"Total registros a normalizar: {to_process} (Filtro: {filter_mode.upper()}, Instancias concurrentes: {workers_count})")

        if to_process == 0:
            state.add_log(f"[INFO] No hay registros pendientes para procesar en [{target_schema}.{target_table}] con filtro '{filter_mode}'.")
            state.finish_run(summary_data={
                "total_records": 0,
                "processed_records": 0,
                "valid_processed_records": 0,
                "observed_records": 0,
                "successful_records": 0,
                "failed_records": 0,
                "ai_records": 0,
                "hybrid_records": 0,
                "heuristic_records": 0,
            })
            return

        summary = pipeline.run(max_records=req.limit, limit=req.limit, filter_mode=filter_mode)

        summary_dict = {
            "total_records": summary.total_records,
            "processed_records": summary.processed_records,
            "valid_processed_records": summary.valid_processed_records,
            "observed_records": summary.observed_records,
            "successful_records": summary.successful_records,
            "failed_records": summary.failed_records,
            "ai_records": summary.ai_records,
            "hybrid_records": summary.hybrid_records,
            "heuristic_records": summary.heuristic_records,
        }
        state.add_log("Pipeline completado exitosamente.")
        state.finish_run(summary_data=summary_dict)

    except Exception as e:
        logger.exception("Error durante la ejecución del pipeline: %s", e)
        state.add_log(f"ERROR: {e}")
        state.finish_run(error=str(e))


@app.post("/api/start")
async def start_pipeline(req: StartPipelineRequest):
    """Inicia la ejecución del pipeline ETL en segundo plano."""
    if state.is_running:
        raise HTTPException(
            status_code=409,
            detail=f"Ya hay un pipeline en ejecución activa en el esquema [{state.active_schema}] tabla [{state.active_table}]."
        )

    state.active_schema = req.schema_name
    state.active_table = req.table_name
    state.num_workers = max(1, min(16, int(req.num_workers or state.num_workers or 4)))

    worker = threading.Thread(
        target=_run_pipeline_worker,
        args=(req,),
        name="ETLWorkerThread",
        daemon=True,
    )
    state.thread = worker
    worker.start()

    return {"status": "STARTED", "message": "Pipeline iniciado correctamente."}


@app.post("/api/stop")
async def stop_pipeline():
    """Detiene la ejecución del pipeline de manera ordenada tras culminar el registro en curso."""
    if not state.is_running or not state.pipeline:
        return {"status": "NOT_RUNNING", "message": "No hay un pipeline en ejecución activa."}

    state.pipeline.request_stop()
    state.stats["status"] = "STOPPING"
    state.broadcast({"type": "stopping", "payload": {"status": "STOPPING", "message": "Detención solicitada."}})
    state.add_log("[INFO] Solicitud de detención manual recibida. Esperando finalización del registro actual...")
    return {"status": "STOPPING", "message": "Detención solicitada. El proceso finalizará en breve."}


@app.post("/api/clear-session")
async def clear_session_endpoint():
    """Limpia los registros acumulados y logs de la sesión activa en el servidor."""
    state.clear_session()
    return {"status": "CLEARED", "message": "Historial de registros de la sesión reiniciado correctamente."}


def generate_excel_report(records: List[Dict[str, Any]]) -> Any:
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Direcciones MPCH"

    headers = [
        "ID Licencia",
        "Dirección Registrada",
        "Tipo Vía",
        "Nombre Vía Homologada",
        "N° Vía",
        "ID Vía Oficial",
        "Tipo Zona",
        "Nombre Zona Homologada",
        "ID Zona Oficial",
        "Manzana",
        "Lote",
        "Sublote",
        "Referencia",
        "Método Normalización",
        "Estado Catastral",
        "Diagnóstico / Observación",
        "Hora de Procesamiento",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="0C2340", end_color="0C2340", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    fill_valid = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    font_valid = Font(name="Calibri", size=11, bold=True, color="065F46")
    fill_obs   = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    font_obs   = Font(name="Calibri", size=11, bold=True, color="92400E")
    fill_err   = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    font_err   = Font(name="Calibri", size=11, bold=True, color="991B1B")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for row_idx, r in enumerate(records, start=2):
        is_err = not r.get("success", True) or r.get("metodo") == "ERROR"
        is_obs = not r.get("es_procesado", False) and not is_err
        estado = "ERROR" if is_err else ("OBSERVADO" if is_obs else "NORMALIZADO")

        row_data = [
            r.get("id_licencia", ""),
            r.get("raw_text", ""),
            r.get("tipo_via_name", ""),
            r.get("nom_via", ""),
            r.get("num_via", ""),
            r.get("id_via", "") or "",
            r.get("tipo_zona_name", ""),
            r.get("nom_zona", ""),
            r.get("id_zona", "") or "",
            r.get("manzana", ""),
            r.get("lote", ""),
            r.get("slote", ""),
            r.get("referencia", ""),
            r.get("metodo", ""),
            estado,
            r.get("observacion", ""),
            r.get("time", ""),
        ]
        ws.append(row_data)

        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx == 15:  # Columna Estado
                if estado == "NORMALIZADO":
                    cell.fill = fill_valid
                    cell.font = font_valid
                elif estado == "OBSERVADO":
                    cell.fill = fill_obs
                    cell.font = font_obs
                else:
                    cell.fill = fill_err
                    cell.font = font_err
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (1, 5, 6, 9):
                cell.alignment = Alignment(horizontal="center")
                if is_even:
                    cell.fill = fill_zebra
            else:
                if is_even:
                    cell.fill = fill_zebra

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_csv_report(records: List[Dict[str, Any]]) -> Any:
    """Genera un archivo CSV con codificación UTF-8 con BOM para compatibilidad directa con Excel en español."""
    import csv
    import io

    output = io.StringIO()
    output.write("\ufeff")  # UTF-8 BOM

    headers = [
        "ID",
        "Direccion_Original",
        "Tipo_Via",
        "Nombre_Via",
        "Numero_Via",
        "ID_Via",
        "Tipo_Zona",
        "Nombre_Zona",
        "ID_Zona",
        "Manzana",
        "Lote",
        "Sublote",
        "Referencia",
        "Metodo_Normalizacion",
        "Estado",
        "Diagnostico_Observacion",
        "Hora_Proceso",
    ]

    writer = csv.writer(output, delimiter=",", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)

    for r in records:
        is_err = not r.get("success", True) or r.get("metodo") == "ERROR"
        is_obs = not r.get("es_procesado", False) and not is_err
        estado = "ERROR" if is_err else ("OBSERVADO" if is_obs else "NORMALIZADO")

        writer.writerow([
            r.get("id_licencia", ""),
            r.get("raw_text", ""),
            r.get("tipo_via_name", ""),
            r.get("nom_via", ""),
            r.get("num_via", ""),
            r.get("id_via", "") or "",
            r.get("tipo_zona_name", ""),
            r.get("nom_zona", ""),
            r.get("id_zona", "") or "",
            r.get("manzana", ""),
            r.get("lote", ""),
            r.get("slote", ""),
            r.get("referencia", ""),
            r.get("metodo", ""),
            estado,
            r.get("observacion", ""),
            r.get("time", ""),
        ])

    buf = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    buf.seek(0)
    return buf


def fetch_db_records_for_export(
    db_svc: DatabaseService,
    schema: Optional[str] = None,
    table: Optional[str] = None,
    scope: str = "processed",  # 'processed', 'observed', 'valid', 'all'
    id_col: Optional[str] = None,
    dir_col: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Consulta la tabla en PostgreSQL con resolución dinámica de esquemas, catálogos y nombres estándar."""
    from sqlalchemy import text as sa_text
    from src.transformers.catalog_matcher import CatalogMatcher
    from src.config.settings import get_settings

    settings = get_settings()
    schema = (schema or state.active_schema or settings.db.schema).strip()
    table = (table or state.active_table or settings.db.table).strip()

    try:
        CatalogMatcher.sync_with_db(db_svc, schema)
    except Exception as e:
        logger.debug("Aviso al sincronizar catálogos para exportación en '%s': %s", schema, e)

    with db_svc.get_session() as session:
        # 1. Comprobar si la tabla existe en el esquema. Si no existe, buscar nombres estándar comunes.
        def check_table_exists(t_name: str) -> bool:
            return bool(session.execute(sa_text("""
                SELECT COUNT(*) 
                FROM information_schema.tables 
                WHERE table_schema = :schema AND table_name = :table;
            """), {"schema": schema, "table": t_name}).scalar())

        if not check_table_exists(table):
            candidates = [
                state.active_table,
                settings.db.table,
                "direcciones_licencias",
                "direcciones_actual",
                f"{table}_actual",
                table.replace("_actual", ""),
                "direcciones",
            ]
            found = False
            for cand in candidates:
                if cand and check_table_exists(cand):
                    logger.info("Tabla '%s' no existe en '%s'. Usando estándar detectado: '%s'", table, schema, cand)
                    table = cand
                    found = True
                    break
            if not found:
                raise ValueError(f"La tabla '{table}' no existe en el esquema '{schema}'.")

        # 2. Inspeccionar columnas existentes en la tabla seleccionada
        col_rows = session.execute(sa_text("""
            SELECT column_name 
            FROM information_schema.columns 
            WHERE table_schema = :schema AND table_name = :table;
        """), {"schema": schema, "table": table}).fetchall()
        existing_cols = {r[0].lower() for r in col_rows}
        actual_cols_map = {r[0].lower(): r[0] for r in col_rows}

        # 3. Identificar columna ID estándar
        target_id_col = None
        for candidate_id in [id_col, settings.db.id_col, "id_licencia", "id", "id_direccion", "codigo"]:
            if candidate_id and candidate_id.lower() in existing_cols:
                target_id_col = actual_cols_map[candidate_id.lower()]
                break
        if not target_id_col:
            target_id_col = col_rows[0][0] if col_rows else "id_licencia"

        # 4. Identificar columna Dirección estándar
        target_dir_col = None
        for candidate_dir in [dir_col, settings.db.dir_col, "emp_direccion", "direccion", "dir", "raw_text", "texto_direccion"]:
            if candidate_dir and candidate_dir.lower() in existing_cols:
                target_dir_col = actual_cols_map[candidate_dir.lower()]
                break

        # 5. Detectar columnas normalizadas en la tabla
        def get_col(conf_name: Optional[str], def_name: str) -> Optional[str]:
            if conf_name and conf_name.lower() in existing_cols:
                return actual_cols_map[conf_name.lower()]
            if def_name and def_name.lower() in existing_cols:
                return actual_cols_map[def_name.lower()]
            return None

        c_id_via       = get_col(settings.db.col_id_via, "id_via")
        c_num_via      = get_col(settings.db.col_num_via, "num_via")
        c_id_zona      = get_col(settings.db.col_id_zona, "id_zona")
        c_manzana      = get_col(settings.db.col_manzana, "manzana")
        c_lote         = get_col(settings.db.col_lote, "lote")
        c_slote        = get_col(settings.db.col_slote, "slote")
        c_referencia   = get_col(settings.db.col_referencia, "referencia")
        c_es_procesado = get_col(settings.db.col_es_procesado, "es_procesado")
        c_observacion  = get_col(settings.db.col_observacion, "observacion")

        # 6. Evaluar scope y filtro WHERE
        scope_lower = str(scope).lower().strip()
        if c_es_procesado:
            if scope_lower in ("observed", "observados"):
                where_clause = f'd."{c_es_procesado}" IS FALSE'
            elif scope_lower in ("valid", "validos", "normalizados"):
                where_clause = f'd."{c_es_procesado}" IS TRUE'
            elif scope_lower in ("processed", "procesados"):
                where_clause = f'd."{c_es_procesado}" IS NOT NULL'
            else:
                where_clause = "TRUE"
        else:
            if scope_lower in ("observed", "observados", "valid", "validos", "normalizados", "processed", "procesados"):
                return []
            where_clause = "TRUE"

        # 7. Inspeccionar tablas maestras en el esquema para resolver JOINs de forma estándar
        vias_table = settings.db.table_vias
        zonas_table = settings.db.table_zonas
        tipo_via_table = settings.db.table_tipo_via
        tipo_zona_table = settings.db.table_tipo_zona

        has_vias = bool(c_id_via and check_table_exists(vias_table))
        has_zonas = bool(c_id_zona and check_table_exists(zonas_table))
        has_tipo_via = check_table_exists(tipo_via_table)
        has_tipo_zona = check_table_exists(tipo_zona_table)

        # Determinar columnas exactas en vias y zonas si existen
        vias_cols = set()
        if has_vias:
            v_col_rows = session.execute(sa_text("""
                SELECT column_name FROM information_schema.columns
                WHERE table_schema = :schema AND table_name = :table;
            """), {"schema": schema, "table": vias_table}).fetchall()
            vias_cols = {r[0].lower() for r in v_col_rows}

        zonas_cols = set()
        if has_zonas:
            z_col_rows = session.execute(sa_text("""
                SELECT column_name FROM information_schema.columns
                WHERE table_schema = :schema AND table_name = :table;
            """), {"schema": schema, "table": zonas_table}).fetchall()
            zonas_cols = {r[0].lower() for r in z_col_rows}

        select_fields = [
            f'd."{target_id_col}" AS id_licencia',
            f'd."{target_dir_col}" AS raw_text' if target_dir_col else "NULL AS raw_text",
            f'd."{c_id_via}" AS id_via' if c_id_via else "NULL AS id_via",
            f'd."{c_num_via}" AS num_via' if c_num_via else "NULL AS num_via",
            f'd."{c_id_zona}" AS id_zona' if c_id_zona else "NULL AS id_zona",
            f'd."{c_manzana}" AS manzana' if c_manzana else "NULL AS manzana",
            f'd."{c_lote}" AS lote' if c_lote else "NULL AS lote",
            f'd."{c_slote}" AS slote' if c_slote else "NULL AS slote",
            f'd."{c_referencia}" AS referencia' if c_referencia else "NULL AS referencia",
            f'd."{c_es_procesado}" AS es_procesado' if c_es_procesado else "NULL AS es_procesado",
            f'd."{c_observacion}" AS observacion' if c_observacion else "NULL AS observacion",
        ]

        join_clauses = []

        # Enriquecimiento Vías
        if has_vias and "id_via" in vias_cols:
            nom_via_expr = 'v."nom_via"' if "nom_via" in vias_cols else "NULL"
            select_fields.append(f"{nom_via_expr} AS nom_via")
            if "id_tipo_via" in vias_cols:
                select_fields.append('v."id_tipo_via" AS id_tipo_via')
                if has_tipo_via:
                    select_fields.append('tv."nombre_tipo_via" AS tipo_via_name')
                    join_clauses.append(f'LEFT JOIN "{schema}"."{vias_table}" v ON d."{c_id_via}" = v."id_via"')
                    join_clauses.append(f'LEFT JOIN "{schema}"."{tipo_via_table}" tv ON v."id_tipo_via" = tv."id_tipo_via"')
                else:
                    select_fields.append('NULL AS tipo_via_name')
                    join_clauses.append(f'LEFT JOIN "{schema}"."{vias_table}" v ON d."{c_id_via}" = v."id_via"')
            elif "tipo_via" in vias_cols:
                select_fields.append('v."tipo_via" AS id_tipo_via')
                select_fields.append('NULL AS tipo_via_name')
                join_clauses.append(f'LEFT JOIN "{schema}"."{vias_table}" v ON d."{c_id_via}" = v."id_via"')
            else:
                select_fields.append('NULL AS id_tipo_via')
                select_fields.append('NULL AS tipo_via_name')
                join_clauses.append(f'LEFT JOIN "{schema}"."{vias_table}" v ON d."{c_id_via}" = v."id_via"')
        else:
            select_fields.extend(["NULL AS nom_via", "NULL AS id_tipo_via", "NULL AS tipo_via_name"])

        # Enriquecimiento Zonas
        if has_zonas and "id_zona" in zonas_cols:
            nom_zona_expr = 'z."nom_zona"' if "nom_zona" in zonas_cols else "NULL"
            select_fields.append(f"{nom_zona_expr} AS nom_zona")
            if "id_tipo_zona" in zonas_cols:
                select_fields.append('z."id_tipo_zona" AS id_tipo_zona')
                if has_tipo_zona:
                    select_fields.append('tz."nombre_tipo_zona" AS tipo_zona_name')
                    join_clauses.append(f'LEFT JOIN "{schema}"."{zonas_table}" z ON d."{c_id_zona}" = z."id_zona"')
                    join_clauses.append(f'LEFT JOIN "{schema}"."{tipo_zona_table}" tz ON z."id_tipo_zona" = tz."id_tipo_zona"')
                else:
                    select_fields.append('NULL AS tipo_zona_name')
                    join_clauses.append(f'LEFT JOIN "{schema}"."{zonas_table}" z ON d."{c_id_zona}" = z."id_zona"')
            elif "tipo_zona" in zonas_cols:
                select_fields.append('z."tipo_zona" AS id_tipo_zona')
                select_fields.append('NULL AS tipo_zona_name')
                join_clauses.append(f'LEFT JOIN "{schema}"."{zonas_table}" z ON d."{c_id_zona}" = z."id_zona"')
            else:
                select_fields.append('NULL AS id_tipo_zona')
                select_fields.append('NULL AS tipo_zona_name')
                join_clauses.append(f'LEFT JOIN "{schema}"."{zonas_table}" z ON d."{c_id_zona}" = z."id_zona"')
        else:
            select_fields.extend(["NULL AS nom_zona", "NULL AS id_tipo_zona", "NULL AS tipo_zona_name"])

        fields_sql = ",\n            ".join(select_fields)
        joins_sql = "\n        ".join(join_clauses)
        query_sql = f"""
            SELECT 
                {fields_sql}
            FROM "{schema}"."{table}" d
            {joins_sql}
            WHERE {where_clause}
            ORDER BY d."{target_id_col}" ASC;
        """

        result_proxy = session.execute(sa_text(query_sql))
        keys = list(result_proxy.keys())
        raw_rows = result_proxy.fetchall()

        records = []
        for raw in raw_rows:
            r = dict(zip(keys, raw))

            id_tipo_via = r.get("id_tipo_via")
            tipo_via_name = r.get("tipo_via_name")
            if not tipo_via_name and id_tipo_via is not None:
                tipo_via_name = CatalogMatcher.get_via_name(id_tipo_via)

            id_tipo_zona = r.get("id_tipo_zona")
            tipo_zona_name = r.get("tipo_zona_name")
            if not tipo_zona_name and id_tipo_zona is not None:
                tipo_zona_name = CatalogMatcher.get_zona_name(id_tipo_zona)

            es_proc = r.get("es_procesado")
            if es_proc is True:
                estado = "NORMALIZADO"
                metodo = "Catastro Oficial"
            elif es_proc is False:
                estado = "OBSERVADO"
                metodo = "Evaluado (Observado)"
            else:
                estado = "PENDIENTE"
                metodo = "Sin Procesar"

            records.append({
                "id_licencia": r.get("id_licencia"),
                "raw_text": r.get("raw_text") or "",
                "id_via": r.get("id_via"),
                "nom_via": r.get("nom_via") or "",
                "tipo_via": id_tipo_via,
                "tipo_via_name": tipo_via_name or "",
                "num_via": r.get("num_via") or "",
                "id_zona": r.get("id_zona"),
                "nom_zona": r.get("nom_zona") or "",
                "tipo_zona": id_tipo_zona,
                "tipo_zona_name": tipo_zona_name or "",
                "manzana": r.get("manzana") or "",
                "lote": r.get("lote") or "",
                "slote": r.get("slote") or "",
                "referencia": r.get("referencia") or "",
                "es_procesado": bool(es_proc) if es_proc is not None else False,
                "observacion": r.get("observacion") or "",
                "estado": estado,
                "metodo": metodo,
                "success": True,
                "time": "-",
            })

    return records


class ExportDbRequest(BaseModel):
    schema_name: Optional[str] = None
    table_name: Optional[str] = None
    scope: str = "processed"  # 'processed', 'observed', 'valid', 'all'
    format: str = "excel"     # 'excel', 'csv'
    id_col: Optional[str] = None
    dir_col: Optional[str] = None


@app.post("/api/export-excel")
async def export_excel_endpoint(payload: Optional[Dict[str, Any]] = None):
    """Genera y descarga un libro de Excel (.xlsx) con formato institucional MPCH desde registros en memoria."""
    from datetime import datetime

    records = []
    if payload and isinstance(payload, dict):
        records = payload.get("records", [])
    if not records:
        records = state.recent_records

    if not records:
        raise HTTPException(status_code=400, detail="No hay registros disponibles para exportar.")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    buf = generate_excel_report(records)
    filename = f"reporte_catastral_mpch_{timestamp}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@app.post("/api/export-csv")
async def export_csv_endpoint(payload: Optional[Dict[str, Any]] = None):
    """Genera y descarga un archivo CSV con formato UTF-8 BOM desde registros en memoria."""
    from datetime import datetime

    records = []
    if payload and isinstance(payload, dict):
        records = payload.get("records", [])
    if not records:
        records = state.recent_records

    if not records:
        raise HTTPException(status_code=400, detail="No hay registros disponibles para exportar.")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    buf = generate_csv_report(records)
    filename = f"reporte_catastral_mpch_{timestamp}.csv"

    return StreamingResponse(
        buf,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@app.api_route("/api/export-db", methods=["GET", "POST"])
async def export_db_endpoint(
    req: Optional[ExportDbRequest] = None,
    schema: Optional[str] = Query(None),
    table: Optional[str] = Query(None),
    scope: str = Query("processed"),
    format: str = Query("excel"),
    id_col: Optional[str] = Query(None),
    dir_col: Optional[str] = Query(None),
):
    """Genera y descarga exportación masiva directamente desde PostgreSQL en Excel o CSV."""
    from datetime import datetime

    db_svc = state.dynamic_db_service or DatabaseService()
    settings = get_settings()

    target_schema = (req.schema_name if req and req.schema_name else None) or schema or state.active_schema or settings.db.schema
    target_table = (req.table_name if req and req.table_name else None) or table or state.active_table or settings.db.table
    target_scope = (req.scope if req and req.scope else None) or scope or "processed"
    target_format = (req.format if req and req.format else None) or format or "excel"
    target_id_col = (req.id_col if req and req.id_col else None) or id_col or settings.db.id_col
    target_dir_col = (req.dir_col if req and req.dir_col else None) or dir_col or settings.db.dir_col

    try:
        records = fetch_db_records_for_export(
            db_svc=db_svc,
            schema=target_schema,
            table=target_table,
            scope=target_scope,
            id_col=target_id_col,
            dir_col=target_dir_col,
        )
    except Exception as e:
        logger.error("Error extrayendo registros para exportar de %s.%s: %s", target_schema, target_table, e)
        raise HTTPException(status_code=400, detail=f"Error consultando registros en base de datos: {str(e)}")

    if not records:
        raise HTTPException(status_code=404, detail=f"No se encontraron registros con filtro '{target_scope}' en {target_schema}.{target_table}.")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    scope_slug = target_scope.lower()

    if target_format.lower() in ("csv", "text/csv"):
        buf = generate_csv_report(records)
        filename = f"reporte_{target_table}_{scope_slug}_{timestamp}.csv"
        media_type = "text/csv; charset=utf-8"
    else:
        buf = generate_excel_report(records)
        filename = f"reporte_{target_table}_{scope_slug}_{timestamp}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return StreamingResponse(
        buf,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )



@app.post("/api/shutdown")
async def shutdown_application():
    """Detiene cualquier pipeline ETL en ejecución y apaga el servicio de la aplicación (equivalente a Ctrl + C con confirmación)."""
    try:
        if state.is_running and state.pipeline:
            state.pipeline.request_stop()
            state.add_log("[APAGADO] Solicitud de apagado de servicio: deteniendo proceso ETL...")
            state.finish_run(error="Servicio apagado por el usuario.")
    except Exception as e:
        logger.warning("Aviso al detener pipeline durante apagado: %s", e)

    state.shutdown()

    def _trigger_ctrl_c():
        time.sleep(0.5)  # Breve lapso para despachar la respuesta HTTP 200 al navegador
        logger.info("[APAGADO] Apagando servicio de aplicación (equivalente a Ctrl + C)...")
        try:
            # Enviar SIGINT (equivalente a Ctrl + C) al proceso de la aplicación
            os.kill(os.getpid(), signal.SIGINT)
        except Exception:
            os._exit(0)
        # Salvaguarda por si el loop de eventos no culminara
        time.sleep(1.2)
        os._exit(0)

    threading.Thread(target=_trigger_ctrl_c, daemon=True).start()

    return {
        "status": "SHUTTING_DOWN",
        "message": "Servicio apagado correctamente.",
    }


@app.get("/api/stream")
async def sse_event_stream():
    """Canal SSE (Server-Sent Events) para transmitir registros y métricas en tiempo real."""
    client_queue: asyncio.Queue = asyncio.Queue(maxsize=500)
    state.register_client(client_queue)

    async def event_generator():
        try:
            # Enviar saludo inicial y estado actual
            yield f"data: {{\"type\": \"connected\", \"payload\": {{\"running\": {str(state.is_running).lower()}}}}}\n\n"

            # Enviar snapshot inicial para sincronización multisesión (dashboard, contadores y registros recientes)
            init_payload = {
                "type": "init",
                "payload": {
                    "is_running": state.is_running,
                    "active_schema": state.active_schema,
                    "active_table": state.active_table,
                    "active_model": state.active_model,
                    "num_workers": state.num_workers,
                    "stats": state.stats,
                    "records": list(state.recent_records)[-50:],
                    "logs": list(state.log_history)[-50:],
                }
            }
            yield f"data: {json.dumps(init_payload, default=str)}\n\n"

            while True:
                try:
                    # Esperar evento con timeout para enviar ping periódico
                    event = await asyncio.wait_for(client_queue.get(), timeout=15.0)
                    if event is None:
                        # Centinela de cierre: finalizar stream ordenadamente
                        break
                    json_data = json.dumps(event, default=str)
                    yield f"data: {json_data}\n\n"
                except asyncio.TimeoutError:
                    yield ": ping\n\n"
        except (asyncio.CancelledError, GeneratorExit):
            pass
        finally:
            state.unregister_client(client_queue)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )
