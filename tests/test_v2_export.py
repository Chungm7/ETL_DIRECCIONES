"""Pruebas unitarias y de integración para el subsistema de exportación relacional V2 (28 columnas y JSON anidado)."""

import csv
import io
import json
import openpyxl
import pytest
from unittest.mock import MagicMock, patch
from fastapi.testclient import TestClient

from src.ui.server import (
    app,
    V2_EXPORT_HEADERS,
    extract_v2_export_row,
    generate_excel_report,
    generate_csv_report,
    generate_json_report,
    fetch_db_records_for_export,
)


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def sample_v2_records():
    return [
        {
            "id_licencia": 1001,
            "raw_text": "CA. SAN JOSE 102 CON AV. LUIS GONZALES 801 MZ B LT 14 DPTO 301 PISO 3 FRENTE AL PARQUE",
            "dire_id": 501,
            "estado": "NORMALIZADO",
            "es_procesado": True,
            "metodo": "IA (patroclo-artesano-7b)",
            "observacion": "Validado exitosamente con catastro oficial",
            "time": "2026-09-25 11:50:00",
            "tipo_zona_name": "URBANIZACION",
            "nom_zona": "SANTA VICTORIA",
            "id_zona": 1,
            "vias": [
                {
                    "via_id": 2860,
                    "via_nombre": "SAN JOSE",
                    "divi_numero": "102",
                    "divi_orden": 1,
                    "tipo_via": 2,
                    "tipo_via_name": "CALLE",
                },
                {
                    "via_id": 2855,
                    "via_nombre": "LUIS GONZALES",
                    "divi_numero": "801",
                    "divi_orden": 2,
                    "tipo_via": 1,
                    "tipo_via_name": "AVENIDA",
                },
            ],
            "componentes": [
                {"codi_id": 1, "codi_nombre": "MANZANA", "diti_nombre": "B"},
                {"codi_id": 2, "codi_nombre": "LOTE", "diti_nombre": "14"},
                {"codi_id": 3, "codi_nombre": "SUBLOTE", "diti_nombre": "A"},
                {"codi_id": 4, "codi_nombre": "PISO", "diti_nombre": "3"},
                {"codi_id": 6, "codi_nombre": "VALLE", "diti_nombre": "CHICLAYO"},
            ],
            "modulos": [
                {"timo_id": 2, "timo_nombre": "DEPARTAMENTO", "ditm_nombre": "301"},
                {"timo_id": 4, "timo_nombre": "STAND", "ditm_nombre": "15"},
            ],
            "dire_referencia": "FRENTE AL PARQUE PRINCIPAL",
        },
        {
            "id_licencia": 1002,
            "raw_text": "CALLE INEXISTENTE 9999",
            "dire_id": None,
            "estado": "OBSERVADO",
            "es_procesado": False,
            "metodo": "Evaluado (Observado)",
            "observacion": "Vía 'INEXISTENTE' no figura en el catálogo maestro de Chiclayo.",
            "time": "2026-09-25 11:51:00",
            "vias": [],
            "componentes": [],
            "modulos": [],
        },
        {
            "id_licencia": 1003,
            "raw_text": "AV BALTA 250",
            "dire_id": 502,
            "es_procesado": True,
            "nom_via": "BALTA",
            "tipo_via_name": "AVENIDA",
            "num_via": "250",
            "id_via": 10,
            "nom_zona": "CENTRO CIVICO",
            "tipo_zona_name": "URBANIZACION",
            "id_zona": 5,
            "manzana": "C",
            "lote": "8",
            "slote": "INT 4",
            "referencia": "AL COSTADO DEL BANCO",
        },
    ]


def test_v2_export_headers_length_and_order():
    """Verifica que las cabeceras contengan exactamente las 28 columnas relacionales V2."""
    assert len(V2_EXPORT_HEADERS) == 28
    assert V2_EXPORT_HEADERS[0] == "ID"
    assert V2_EXPORT_HEADERS[1] == "ID_Direccion"
    assert V2_EXPORT_HEADERS[2] == "Estado"
    assert V2_EXPORT_HEADERS[3] == "Direccion_Original"
    assert V2_EXPORT_HEADERS[4] == "Via_Principal_Tipo"
    assert V2_EXPORT_HEADERS[5] == "Via_Principal_Nombre"
    assert V2_EXPORT_HEADERS[6] == "Via_Principal_Numero"
    assert V2_EXPORT_HEADERS[7] == "Via_Principal_ID"
    assert V2_EXPORT_HEADERS[8] == "Via_Secundaria_Tipo"
    assert V2_EXPORT_HEADERS[9] == "Via_Secundaria_Nombre"
    assert V2_EXPORT_HEADERS[10] == "Via_Secundaria_Numero"
    assert V2_EXPORT_HEADERS[11] == "Via_Secundaria_ID"
    assert V2_EXPORT_HEADERS[12] == "Todas_Las_Vias"
    assert V2_EXPORT_HEADERS[13] == "Tipo_Zona"
    assert V2_EXPORT_HEADERS[14] == "Nombre_Zona"
    assert V2_EXPORT_HEADERS[15] == "ID_Zona"
    assert V2_EXPORT_HEADERS[16] == "Manzana"
    assert V2_EXPORT_HEADERS[17] == "Lote"
    assert V2_EXPORT_HEADERS[18] == "Sublote"
    assert V2_EXPORT_HEADERS[19] == "Piso"
    assert V2_EXPORT_HEADERS[20] == "Otros_Componentes"
    assert V2_EXPORT_HEADERS[21] == "Tipo_Modulo"
    assert V2_EXPORT_HEADERS[22] == "Numero_Modulo"
    assert V2_EXPORT_HEADERS[23] == "Todos_Los_Modulos"
    assert V2_EXPORT_HEADERS[24] == "Referencia"
    assert V2_EXPORT_HEADERS[25] == "Metodo_Normalizacion"
    assert V2_EXPORT_HEADERS[26] == "Diagnostico_Observacion"
    assert V2_EXPORT_HEADERS[27] == "Fecha_Hora_Proceso"


def test_extract_v2_export_row_multi_via(sample_v2_records):
    """Verifica la extracción y desglose limpio de multi-vías (Vía Principal y Vía Secundaria)."""
    row = extract_v2_export_row(sample_v2_records[0])
    assert row["ID"] == 1001
    assert row["ID_Direccion"] == 501
    assert row["Estado"] == "NORMALIZADO"
    assert row["Via_Principal_Tipo"] == "CALLE"
    assert row["Via_Principal_Nombre"] == "SAN JOSE"
    assert row["Via_Principal_Numero"] == "102"
    assert row["Via_Principal_ID"] == 2860
    assert row["Via_Secundaria_Tipo"] == "AVENIDA"
    assert row["Via_Secundaria_Nombre"] == "LUIS GONZALES"
    assert row["Via_Secundaria_Numero"] == "801"
    assert row["Via_Secundaria_ID"] == 2855
    assert "CALLE SAN JOSE N° 102 con AVENIDA LUIS GONZALES N° 801" in row["Todas_Las_Vias"]
    assert row["Tipo_Zona"] == "URBANIZACION"
    assert row["Nombre_Zona"] == "SANTA VICTORIA"
    assert row["ID_Zona"] == 1
    assert row["Manzana"] == "B"
    assert row["Lote"] == "14"
    assert row["Sublote"] == "A"
    assert row["Piso"] == "3"
    assert "VALLE: CHICLAYO" in row["Otros_Componentes"]
    assert row["Tipo_Modulo"] == "DEPARTAMENTO"
    assert row["Numero_Modulo"] == "301"
    assert "DEPARTAMENTO 301, STAND 15" in row["Todos_Los_Modulos"]
    assert row["Referencia"] == "FRENTE AL PARQUE PRINCIPAL"


def test_extract_v2_export_row_observed(sample_v2_records):
    """Verifica que un registro observado conserve su diagnóstico y deje en blanco campos relacionales."""
    row = extract_v2_export_row(sample_v2_records[1])
    assert row["ID"] == 1002
    assert row["ID_Direccion"] == ""
    assert row["Estado"] == "OBSERVADO"
    assert row["Via_Principal_Nombre"] == ""
    assert "no figura en el catálogo maestro" in row["Diagnostico_Observacion"]


def test_extract_v2_export_row_legacy_fallback(sample_v2_records):
    """Verifica la retrocompatibilidad cuando el registro solo tiene campos planos V1."""
    row = extract_v2_export_row(sample_v2_records[2])
    assert row["ID"] == 1003
    assert row["ID_Direccion"] == 502
    assert row["Estado"] == "NORMALIZADO"
    assert row["Via_Principal_Tipo"] == "AVENIDA"
    assert row["Via_Principal_Nombre"] == "BALTA"
    assert row["Via_Principal_Numero"] == "250"
    assert row["Via_Principal_ID"] == 10
    assert row["Via_Secundaria_Nombre"] == ""
    assert "AVENIDA BALTA N° 250" in row["Todas_Las_Vias"]
    assert row["Tipo_Modulo"] == "INT"
    assert row["Numero_Modulo"] == "4"
    assert row["Referencia"] == "AL COSTADO DEL BANCO"


def test_generate_excel_report_structure(sample_v2_records):
    """Verifica que generate_excel_report cree el archivo binario con 28 columnas y estilos."""
    buf = generate_excel_report(sample_v2_records)
    assert isinstance(buf, io.BytesIO)
    content = buf.getvalue()
    assert len(content) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws.title == "Normalización Catastral V2"
    assert ws.max_column == 28
    assert ws.max_row == 4  # 1 fila de cabecera + 3 registros

    # Verificar cabeceras
    header_cells = [ws.cell(row=1, column=col).value for col in range(1, 29)]
    assert header_cells == V2_EXPORT_HEADERS

    # Verificar primera fila
    row1 = [ws.cell(row=2, column=col).value for col in range(1, 29)]
    assert row1[0] == 1001
    assert row1[1] == 501
    assert row1[2] == "NORMALIZADO"
    assert row1[4] == "CALLE"
    assert row1[5] == "SAN JOSE"
    assert row1[8] == "AVENIDA"
    assert row1[9] == "LUIS GONZALES"


def test_generate_csv_report_utf8_bom(sample_v2_records):
    """Verifica que generate_csv_report cree un archivo CSV con BOM UTF-8 y 28 columnas."""
    buf = generate_csv_report(sample_v2_records)
    assert isinstance(buf, io.BytesIO)
    raw_bytes = buf.getvalue()
    assert raw_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM

    content = raw_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    assert len(rows) == 4  # Cabecera + 3 filas
    assert rows[0] == V2_EXPORT_HEADERS

    # Fila 1 (Multi-vía)
    assert rows[1][0] == "1001"
    assert rows[1][1] == "501"
    assert rows[1][2] == "NORMALIZADO"
    assert rows[1][5] == "SAN JOSE"
    assert rows[1][9] == "LUIS GONZALES"
    assert "CALLE SAN JOSE N° 102 con AVENIDA LUIS GONZALES N° 801" in rows[1][12]

    # Fila 2 (Observado)
    assert rows[2][0] == "1002"
    assert rows[2][2] == "OBSERVADO"


def test_generate_json_report_hierarchy(sample_v2_records):
    """Verifica que generate_json_report genere la jerarquía institucional anidada."""
    buf = generate_json_report(sample_v2_records)
    assert isinstance(buf, io.BytesIO)
    data = json.loads(buf.getvalue().decode("utf-8"))

    assert "metadata" in data
    assert data["metadata"]["version"] == "2.0"
    assert data["metadata"]["total_registros"] == 3

    dirs = data["direcciones"]
    assert len(dirs) == 3

    # Registro 1 (Multi-vía con anidamiento completo)
    d1 = dirs[0]
    assert d1["id"] == 1001
    assert d1["dire_id"] == 501
    assert d1["estado"] == "NORMALIZADO"
    assert d1["zona"]["nombre_zona"] == "SANTA VICTORIA"
    assert d1["zona"]["tipo_zona"] == "URBANIZACION"
    assert len(d1["vias"]) == 2
    assert d1["vias"][0]["orden"] == 1
    assert d1["vias"][0]["nombre_via"] == "SAN JOSE"
    assert d1["vias"][1]["orden"] == 2
    assert d1["vias"][1]["nombre_via"] == "LUIS GONZALES"
    assert len(d1["componentes"]) == 5
    assert d1["catastro"]["manzana"] == "B"
    assert d1["catastro"]["piso"] == "3"
    assert d1["piso"] == "3"
    assert len(d1["modulos"]) == 2
    assert d1["modulos"][0]["tipo"] == "DEPARTAMENTO"
    assert d1["modulos"][0]["numero"] == "301"

    # Retrocompatibilidad
    assert d1["nom_via"] == "SAN JOSE"
    assert d1["es_procesado"] is True


def test_endpoints_export_excel_csv_json(client, sample_v2_records):
    """Verifica que los tres endpoints POST /api/export-* funcionen con registros en payload."""
    payload = {"records": sample_v2_records}

    # Excel
    res_xlsx = client.post("/api/export-excel", json=payload)
    assert res_xlsx.status_code == 200
    assert "application/vnd.openxmlformats" in res_xlsx.headers["content-type"]
    assert len(res_xlsx.content) > 1000

    # CSV
    res_csv = client.post("/api/export-csv", json=payload)
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    csv_text = res_csv.content.decode("utf-8-sig")
    assert "ID,ID_Direccion,Estado,Direccion_Original" in csv_text
    assert "1001" in csv_text

    # JSON
    res_json = client.post("/api/export-json", json=payload)
    assert res_json.status_code == 200
    assert "application/json" in res_json.headers["content-type"]
    json_data = res_json.json()
    assert json_data["metadata"]["version"] == "2.0"
    assert len(json_data["direcciones"]) == 3


def test_api_export_db_json_format(client, sample_v2_records):
    """Verifica que /api/export-db soporte format=json retornando la estructura V2."""
    with patch("src.ui.server.fetch_db_records_for_export", return_value=sample_v2_records):
        res = client.get("/api/export-db?schema=public&table=direcciones_actual&scope=processed&format=json")
        assert res.status_code == 200
        assert "application/json" in res.headers["content-type"]
        data = res.json()
        assert data["metadata"]["version"] == "2.0"
        assert len(data["direcciones"]) == 3
        assert data["direcciones"][0]["zona"]["nombre_zona"] == "SANTA VICTORIA"
